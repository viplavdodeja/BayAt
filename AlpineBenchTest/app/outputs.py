"""Verified analytics export; records.json remains authoritative."""
import math
import os
import tempfile
from datetime import date, datetime
from pathlib import Path
from .analytics import run_rows


def export_analytics(data, steps, job_directory):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter
    headers, rows = run_rows(data, steps)
    directory = Path(job_directory) / 'Analysis'
    directory.mkdir(parents=True, exist_ok=True)
    destination = directory / (data['job'] + '_analysis.xlsx')
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Test runs'
    sheet.freeze_panes = 'D2'
    for row in [headers, *rows]:
        sheet.append(row)
    for row in sheet.iter_rows():
        for cell in row:
            value = cell.value
            if isinstance(value, str):
                if len(value) > 32767:
                    raise ValueError(f'Text exceeds Excel cell limit at {cell.coordinate}.')
                cell.data_type = 's'
                cell.number_format = '@'
            elif isinstance(value, datetime):
                cell.number_format = 'yyyy-mm-dd hh:mm:ss.000'
            elif isinstance(value, date):
                cell.number_format = 'yyyy-mm-dd'
            elif isinstance(value, (float, int)) and not isinstance(value, bool):
                cell.number_format = '0.00000' if headers[cell.column - 1].endswith('_decay') else '0.###############'
            if cell.row == 1:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='234E70')
                cell.alignment = Alignment(wrap_text=True, vertical='center')
    sheet.row_dimensions[1].height = 72
    for index, header in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(index)].width = 38 if header in ('Run ID', 'Notes', 'Previously failed steps') else 24
    if rows:
        table = Table(displayName='TestRuns', ref=f'A1:{get_column_letter(len(headers))}{len(rows)+1}')
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        sheet.add_table(table)
    descriptor, temporary = tempfile.mkstemp(suffix='.xlsx', dir=directory)
    os.close(descriptor)
    try:
        workbook.save(temporary)
        check = load_workbook(temporary, read_only=True, data_only=False)
        try:
            restored = check['Test runs']
            expected = [headers, *rows]
            if restored.max_row != len(expected) or restored.max_column != len(headers):
                raise ValueError('Excel verification failed: row/column count differs.')
            for actual_row, expected_row in zip(restored.iter_rows(), expected):
                for cell, value in zip(actual_row, expected_row):
                    actual = cell.value
                    if value == '':
                        value = None
                    if isinstance(value, date) and not isinstance(value, datetime):
                        actual = actual.date() if isinstance(actual, datetime) else actual
                    elif isinstance(value, datetime) and isinstance(actual, datetime):
                        if abs((actual-value).total_seconds()) <= .001:
                            actual = value
                    elif isinstance(value, float) and isinstance(actual, (int, float)):
                        if math.isclose(actual, value, rel_tol=1e-15, abs_tol=0):
                            actual = value
                    if actual != value or (isinstance(value, bool) and cell.data_type != 'b') or (isinstance(value, str) and cell.data_type != 's'):
                        raise ValueError(f'Excel verification failed at {cell.coordinate}.')
        finally:
            check.close()
        os.replace(temporary, destination)
    finally:
        workbook.close()
        if os.path.exists(temporary):
            os.unlink(temporary)
    return destination


def main():
    import argparse
    from .storage import Store, load_config
    parser = argparse.ArgumentParser(description='Regenerate analytics from saved records.')
    parser.add_argument('job', nargs='?', help='Job number; omit for all saved jobs.')
    args = parser.parse_args()
    store = Store()
    jobs = [args.job] if args.job else [p.parent.name for p in store.root.glob('*/records.json')]
    for job in jobs:
        print(export_analytics(store.load(job), load_config('test_steps.json'), store.job_path(job)))


if __name__ == '__main__':
    main()
