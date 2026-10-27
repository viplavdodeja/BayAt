import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from app.storage import Store, load_config
from app.outputs import export_analytics


class OutputTests(unittest.TestCase):
    def test_types_formula_text_and_stable_rows(self):
        with tempfile.TemporaryDirectory() as destination:
            store = Store(destination)
            data = store.load("0001")
            run = store.new_run(data, "000012345678901234567", "2026-10-11")
            run["notes"] = ["=literal note"]
            run["attempts"] = [{"step": "cda1", "title": "CDA 1", "passed": True,
                                "values": {"QDF6_initial": 10.123456789, "QDF6_pass": True}}]
            store.save(data)
            self.assertIsNone(store.export_error)
            path = Path(destination) / "0001/Analysis/0001_analysis.xlsx"
            with_workbook = load_workbook(path)
            sheet = with_workbook["Test runs"]
            self.assertEqual(sheet.max_row, 2)
            self.assertEqual(sheet['A2'].value, "0001")
            self.assertEqual(sheet['C2'].value, run["serial"])
            self.assertEqual(sheet['L2'].value, "=literal note")
            self.assertEqual(sheet['L2'].data_type, "s")
            self.assertEqual(sheet['D2'].value.date().isoformat(), "2026-10-11")
            self.assertIn("TestRuns", sheet.tables)
            with_workbook.close()

    def test_export_failure_preserves_records_and_old_workbook(self):
        with tempfile.TemporaryDirectory() as destination:
            store = Store(destination)
            data = store.load("job")
            run = store.new_run(data, "serial", "2026-10-11")
            path = Path(destination) / "job/Analysis/job_analysis.xlsx"
            old = path.read_bytes()
            run["notes"].append("Saved despite locked workbook")
            import os
            original = os.replace
            def replace(source, target):
                if str(target).endswith('.xlsx'):
                    raise PermissionError("Workbook locked")
                return original(source, target)
            with patch("app.outputs.os.replace", side_effect=replace):
                store.save(data)
            self.assertIn("locked", store.export_error)
            self.assertEqual(store.load("job")["runs"][0]["notes"], run["notes"])
            self.assertEqual(path.read_bytes(), old)

    def test_oversize_notes_fail_instead_of_truncating(self):
        with tempfile.TemporaryDirectory() as destination:
            store = Store(destination)
            data = store.load("job")
            run = store.new_run(data, "serial", "2026-10-11")
            run["notes"] = ["x" * 32768]
            with self.assertRaises(ValueError):
                export_analytics(data, load_config("test_steps.json"), store.job_path("job"))
